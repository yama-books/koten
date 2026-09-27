"""Build conj/data/adjectival-noun-glosses.json from the audited v0.6 workbook.

Source (audited, used verbatim):
  形容動詞_現代仮名遣い・基本義・用例補足ストック_v0.6_adjv001-120.xlsx
    sheet `runtime候補`  : id, modernKana, basicGloss, contextGloss, displayGloss
    sheet `監査ストック` : ID, ..., 作品 (work, used only for the §8-2 cross-check)

Design: conj/DESIGN_GLOSS_LAYOUT_2026-09-28.md §6.
  * records[]       : one per v0.6 row (120). contextNote is the inner text of
                      「（この用例では…）」 in displayGloss (null if absent). The
                      contextGloss column is an audit note and is NOT exported.
  * lemmaReadings[] : one per heading the app displays (116 = runtime 115 +
                      built-in 百人一首 `itadura`). The modern reading comes from the
                      audited v0.6 modernKana of the displayed example (stem +
                      なり/たり); the ruby segments are found by aligning the heading
                      with that reading using 歴史的仮名遣い rule units only.
  * itemGlosses[]   : glosses for built-in items that have no v0.6 example
                      (itadura: basicGloss only, §1.3 / §9-3 decision).

The script stops (exit 1) if any check fails: 120 ids without gaps/duplicates,
displayGloss == basicGloss (+ （この用例では contextNote）), the displayed
example's work == v0.6 work for every displayed heading, and every heading
aligns with its modern reading. Hand overrides (none needed at the time of
writing) are listed in HAND_OVERRIDES and printed when used.

Run:  python conj/audit/build_adjv_glosses.py [path/to/v0.6.xlsx]
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
CONJ = HERE.parent
DATA = CONJ / "data"
OUT = DATA / "adjectival-noun-glosses.json"
DEFAULT_XLSX = Path.home() / "Downloads" / "形容動詞_現代仮名遣い・基本義・用例補足ストック_v0.6_adjv001-120.xlsx"
SOURCE_LABEL = "形容動詞_現代仮名遣い・基本義・用例補足ストック_v0.6_adjv001-120.xlsx#runtime候補"

# Built-in (index.html) 形容動詞 items that are not runtime lemmas.
# itadura = 小倉百人一首 九番「いたづらに」. It replaces runtime adjv-lemma-011
# (same heading), whose v0.6 row is adjv-019 (土佐日記). §1.3 decision: show only the
# basic gloss of adjv-019; its context note (暇だ) belongs to the 土佐日記 example.
BUILTIN_ITEMS = [
    {"itemId": "itadura", "heading": "いたづらなり", "readingFrom": "adjv-019", "glossFrom": "adjv-019",
     "work": "小倉百人一首"},
]

# Heading -> ruby segments, for cases the rule aligner cannot decide. Empty on purpose;
# anything added here is printed on every run.
HAND_OVERRIDES: dict[str, list[list[str]]] = {}

CONTEXT_RE = re.compile(r"^(.*)（この用例では(.+)）$")

# --- 歴史的仮名遣い -> 現代仮名遣い rule units ---------------------------------------
A_TO_O = dict(zip("かさたなはまやらわがざだばぱあ", "こそとのほもよろおごぞどぼぽお"))
E_TO_I_YOU = {"け": "きょう", "せ": "しょう", "て": "ちょう", "ね": "にょう", "へ": "ひょう",
              "め": "みょう", "れ": "りょう", "げ": "ぎょう", "ぜ": "じょう", "で": "じょう",
              "べ": "びょう", "ぺ": "ぴょう", "え": "よう"}
I_TO_YUU = {"き": "きゅう", "し": "しゅう", "ち": "ちゅう", "に": "にゅう", "ひ": "ひゅう",
            "み": "みゅう", "り": "りゅう", "ぎ": "ぎゅう", "じ": "じゅう", "び": "びゅう", "い": "ゆう"}
YOU_SMALL = {"き": "きょう", "し": "しょう", "ち": "ちょう", "に": "にょう", "ひ": "ひょう",
             "み": "みょう", "り": "りょう", "ぎ": "ぎょう", "じ": "じょう", "び": "びょう"}


def rule_units() -> list[tuple[str, str, str, bool]]:
    """(historical, modern, rule name, word-internal only)."""
    u: list[tuple[str, str, str, bool]] = []
    for k, v in YOU_SMALL.items():
        u.append((k + "やう", v, "iyau→yō", False))
    u.append(("くわう", "こう", "kwau→kō", False))
    u.append(("ぐわう", "ごう", "gwau→gō", False))
    for k, v in A_TO_O.items():
        u.append((k + "う", v + "う", "au→ō", False))
        u.append((k + "ふ", v + "う", "au(ふ)→ō", False))
    for k, v in E_TO_I_YOU.items():
        u.append((k + "う", v, "eu→yō", False))
        u.append((k + "ふ", v, "eu(ふ)→yō", False))
    for k, v in I_TO_YUU.items():
        u.append((k + "う", v, "iu→yū", False))
        u.append((k + "ふ", v, "iu(ふ)→yū", False))
    u += [("くわ", "か", "kwa→ka", False), ("ぐわ", "が", "gwa→ga", False)]
    u += [(h, m, "語中ハ行→ワ行", True) for h, m in zip("はひふへほ", "わいうえお")]
    u += [("ゐ", "い", "ゐ→い", False), ("ゑ", "え", "ゑ→え", False), ("を", "お", "を→お", False)]
    u += [("ぢ", "じ", "ぢ→じ", False), ("づ", "ず", "づ→ず", False)]
    u += [("む", "ん", "語中む→ん", True), ("つ", "っ", "促音", True)]
    # longest historical unit first, so au/eu/iu groups win over single letters
    return sorted(u, key=lambda x: -len(x[0]))


RULES = rule_units()


def align(heading: str, modern: str):
    """Return ruby segments [[base], [base, rt], ...] or None. Identity is tried last
    at each position so rule units are kept whole (やう→よう, not や→よ)."""
    memo: dict[tuple[int, int], list | None] = {}

    def go(i: int, j: int):
        if (i, j) in memo:
            return memo[(i, j)]
        if i == len(heading):
            res = [] if j == len(modern) else None
            memo[(i, j)] = res
            return res
        res = None
        for h, m, _name, internal in RULES:
            if internal and i == 0:
                continue
            if heading.startswith(h, i) and modern.startswith(m, j):
                rest = go(i + len(h), j + len(m))
                if rest is not None:
                    res = [[h, m]] + rest
                    break
        if res is None and j < len(modern) and heading[i] == modern[j]:
            rest = go(i + 1, j + 1)
            if rest is not None:
                res = [[heading[i]]] + rest
        memo[(i, j)] = res
        return res

    segs = go(0, 0)
    if segs is None:
        return None
    merged: list[list[str]] = []
    for s in segs:
        if len(s) == 1 and merged and len(merged[-1]) == 1:
            merged[-1][0] += s[0]
        else:
            merged.append(list(s))
    return merged


ENDINGS = sorted(["なり", "なる", "なれ", "なら", "に", "たり", "たる", "たれ", "たら", "と"], key=len, reverse=True)


def modern_stem(modern_form: str) -> str:
    for e in ENDINGS:
        if modern_form.endswith(e) and len(modern_form) > len(e):
            return modern_form[: -len(e)]
    raise ValueError("no inflectional ending in " + modern_form)


def read_workbook(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rt = wb["runtime候補"]
    rows = list(rt.iter_rows(values_only=True))
    head = [str(h).strip() for h in rows[0]]
    want = ["id", "modernKana", "basicGloss", "contextGloss", "displayGloss"]
    if head[: len(want)] != want:
        raise SystemExit(f"runtime候補 header mismatch: {head}")
    runtime = [dict(zip(want, r[: len(want)])) for r in rows[1:] if r and r[0]]
    au = wb["監査ストック"]
    arows = list(au.iter_rows(values_only=True))
    ahead = [str(h).strip() if h is not None else "" for h in arows[0]]
    id_col, work_col = ahead.index("ID"), ahead.index("作品")
    works = {r[id_col]: r[work_col] for r in arows[1:] if r and r[id_col]}
    return runtime, works


def text(v):
    return None if v is None else str(v).strip()


def displayed_examples():
    """Replicates conj/adjv-runtime-adapter.js: public example first (4 gates), else the
    first CHJ quotation for the lemma. Returns lemmaId -> (exampleId, work)."""
    pub = json.loads((DATA / "adjectival-noun-public-examples.json").read_text(encoding="utf-8"))
    chj = json.loads((DATA / "adjectival-noun-chj-quotations.json").read_text(encoding="utf-8"))
    out: dict[str, tuple[str, str]] = {}
    for r in pub["records"]:
        ok = (r.get("exampleEnabledPublic") is True and r.get("rightsVerified") is True
              and r.get("targetVerified") is True and r.get("excerptReviewed") is True
              and r.get("example") and r.get("publicTarget") and r["publicTarget"] in r["example"])
        if ok and r["lemmaId"] not in out:
            out[r["lemmaId"]] = (r["sourceExampleId"], r["work"])
    chj_first: dict[str, tuple[str, str]] = {}
    for r in chj["records"]:
        if not r.get("excerpt") or not r.get("target") or r["target"] not in r["excerpt"]:
            continue
        chj_first.setdefault(r["lemmaId"], (r["id"], r["work"]))
    for k, v in chj_first.items():
        out.setdefault(k, v)
    return out


def headings_from_app():
    """lemmaId -> displayed heading, mirroring lemmaHeadingText() in conj/index.html."""
    html = (CONJ / "index.html").read_text(encoding="utf-8")
    m = re.search(r"const HISTORICAL_KANA_HEADINGS=\{([\s\S]*?)\};", html)
    kana = dict(re.findall(r'"([^"]+)":"([^"]+)"', m.group(1)))
    builtin_adjv = set(re.findall(r'\{id:"[^"]+",pos:"adjv",label:"形容動詞",lemma:"([^"]+)"', html))
    pool = json.loads((DATA / "adjectival-noun-lemma-pool.json").read_text(encoding="utf-8"))["lemmas"]
    ann = {a["id"]: a for a in json.loads((DATA / "adjectival-noun-lexical-annotations.json")
                                           .read_text(encoding="utf-8"))["annotations"]}
    out = {}
    for lemma in pool:
        stems = [ann.get(i, {}).get("displayLemma") for i in lemma["exampleIds"]]
        stable = stems[0] if stems and all(s and s == stems[0] for s in stems) else None
        suffix = "たり" if lemma["paradigmId"] == "adjv-tari" else "なり"
        display = stable + suffix if stable else lemma["lemma"]
        if display in builtin_adjv:
            continue  # replaced by a built-in item with the same heading (itadura)
        out[lemma["id"]] = {"heading": kana.get(display, display), "exampleIds": lemma["exampleIds"]}
    return out


def main(argv: list[str]) -> int:
    xlsx = Path(argv[1]) if len(argv) > 1 else DEFAULT_XLSX
    runtime, works = read_workbook(xlsx)
    errors: list[str] = []

    ids = [r["id"] for r in runtime]
    expected = [f"adjv-{n:03d}" for n in range(1, 121)]
    if ids != expected:
        errors.append(f"runtime候補 ids are not adjv-001..120 in order (n={len(ids)})")

    records = []
    by_id = {}
    for r in runtime:
        basic, display = text(r["basicGloss"]), text(r["displayGloss"])
        m = CONTEXT_RE.match(display)
        note = m.group(2) if m else None
        rebuilt = basic + ("（この用例では" + note + "）" if note else "")
        if rebuilt != display:
            errors.append(f"{r['id']}: displayGloss is not basicGloss(+（この用例では…）): {display}")
        work = text(works.get(r["id"]))
        if not work:
            errors.append(f"{r['id']}: work missing in 監査ストック")
        rec = {"id": r["id"], "work": work, "modernKana": text(r["modernKana"]), "basicGloss": basic,
               "contextNote": note, "displayGloss": display}
        records.append(rec)
        by_id[r["id"]] = rec

    shown = displayed_examples()
    headings = headings_from_app()
    readings = []
    used_overrides = []

    def reading(item_id, lemma_id, heading, reading_from):
        suffix = heading[-2:]
        modern = modern_stem(by_id[reading_from]["modernKana"]) + suffix
        if heading in HAND_OVERRIDES:
            ruby = HAND_OVERRIDES[heading]
            used_overrides.append(heading)
        else:
            ruby = align(heading, modern)
        if ruby is None:
            errors.append(f"{item_id}: cannot align {heading} -> {modern} ({reading_from})")
            return
        if "".join(s[0] for s in ruby) != heading or "".join(s[-1] for s in ruby) != modern:
            errors.append(f"{item_id}: ruby does not rebuild {heading} / {modern}")
        if any(len(s) == 2 for s in ruby) != (heading != modern):
            errors.append(f"{item_id}: ruby presence mismatch for {heading}")
        readings.append({"itemId": item_id, "lemmaId": lemma_id, "heading": heading,
                         "modernKana": modern, "ruby": ruby, "readingFrom": reading_from})

    for b in BUILTIN_ITEMS:
        reading(b["itemId"], "adjv-lemma-011", b["heading"], b["readingFrom"])
    for lemma_id, h in headings.items():
        ex = shown.get(lemma_id)
        if not ex:
            errors.append(f"{lemma_id}: no displayed example")
            continue
        ex_id, ex_work = ex
        if ex_id not in by_id:
            errors.append(f"{lemma_id}: displayed example {ex_id} is not in v0.6")
            continue
        if by_id[ex_id]["work"] != ex_work:
            errors.append(f"{lemma_id}: work mismatch {ex_id} app={ex_work} v0.6={by_id[ex_id]['work']}")
        reading(lemma_id, lemma_id, h["heading"], ex_id)

    item_glosses = []
    for b in BUILTIN_ITEMS:
        src = by_id[b["glossFrom"]]
        item_glosses.append({"itemId": b["itemId"], "work": b["work"], "basicGloss": src["basicGloss"],
                             "contextNote": None, "glossFrom": b["glossFrom"],
                             "note": "百人一首九番「いたづらに」＝むなしく・無駄に。基本義のみ表示し、"
                                     "adjv-019（土佐日記）の用例補足は付けない（DESIGN_GLOSS_LAYOUT §1.3）。"})

    if used_overrides:
        print("hand overrides used:", ", ".join(used_overrides))
    if errors:
        print("\n".join("ERROR " + e for e in errors), file=sys.stderr)
        return 1

    out = {
        "schemaVersion": "1.0",
        "source": SOURCE_LABEL,
        "generatedBy": "conj/audit/build_adjv_glosses.py",
        "policy": "records は v0.6 の監査値をそのまま使う。contextNote は displayGloss の「（この用例では…）」の中身。"
                  "画面は表示中の用例 ID（exampleId）で record を引き、作品が一致しないときは表示しない。",
        "records": records,
        "itemGlosses": item_glosses,
        "lemmaReadings": readings,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    diff = [r for r in readings if r["heading"] != r["modernKana"]]
    print(f"wrote {OUT.relative_to(CONJ.parent)}: {len(records)} records, {len(readings)} headings, "
          f"{len(diff)} with ruby ({len({r['heading'] for r in diff})} distinct)")
    for r in diff:
        print("  ", r["itemId"], "".join(f"[{s[0]}→{s[1]}]" if len(s) == 2 else s[0] for s in r["ruby"]))
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main(sys.argv))
